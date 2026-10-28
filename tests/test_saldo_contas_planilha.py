from __future__ import annotations

import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from backend.relatorios.saldo_contas import (
    CONTAS_APLICACAO_MODELO,
    analisar_planilha_saldo_contas,
    detectar_modelo_saldo_contas,
    normalizar_modelo_saldo_contas,
    preencher_planilha_saldo_contas,
    planejar_consultas_saldo_contas,
)


class SaldoContasPlanilhaTests(unittest.TestCase):
    def _criar_planilha(self, pasta: Path, *, alterar_primeira_conta=False, vazia=False):
        caminho = pasta / "saldo-contas.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SET.26"
        sheet["C1"] = date(2026, 9, 1)
        sheet["C19"] = "APLIC. MOVIMENTO"
        sheet["D19"] = "Individuais"

        total = 0

        for indice, conta in enumerate(CONTAS_APLICACAO_MODELO, start=20):
            sheet.cell(indice, 3).value = (
                "99999999" if alterar_primeira_conta and indice == 20 else int(conta)
            )
            if not vazia:
                valor = indice * 10.25
                sheet.cell(indice, 4).value = valor
                total += valor

        sheet["B59"] = "=D59+D60"
        sheet["C61"] = total if not vazia else None
        sheet["D61"] = "=SUM(D20:D60)"
        sheet["J5"] = "=D61-J6"
        sheet["J6"] = "=B59"
        sheet["J7"] = "=J3+J4+J5+J6"
        sheet["D20"].fill = PatternFill(fill_type="solid", fgColor="FFF200")
        sheet["D20"].number_format = '#,##0.00;[Red]-#,##0.00;"-"'
        workbook.save(caminho)
        workbook.close()
        return caminho

    def _criar_planilha_dinamica(self, pasta: Path):
        caminho = pasta / "modelo-dinamico.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OUT.26"
        sheet["A1"] = date(2026, 10, 1)
        sheet["A3"] = "ABACATE MOVIMENTO"
        sheet["A4"] = 11311001
        sheet["A5"] = 11311002
        sheet["A4"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
        sheet["A6"] = 0
        sheet["B6"] = "=SUM(B4:B5)"
        workbook.save(caminho)
        workbook.close()
        return caminho

    def test_detecta_grupo_novo_e_modos_pela_cor(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha_dinamica(Path(temporario))
            deteccao = detectar_modelo_saldo_contas(caminho, "OUT.26")

        self.assertTrue(deteccao["ok"])
        modelo = deteccao["modelo"]
        self.assertEqual(modelo["aba"], "OUT.26")
        self.assertEqual(modelo["celula_periodo"], "A1")
        self.assertEqual(len(modelo["grupos"]), 1)
        grupo = modelo["grupos"][0]
        self.assertEqual(grupo["nome"], "ABACATE MOVIMENTO")
        self.assertEqual(grupo["celula_total"], "A6")
        self.assertEqual([item["conta"] for item in grupo["contas"]], ["11311001", "11311002"])
        self.assertEqual([item["modo"] for item in grupo["contas"]], ["individual", "lote"])
        self.assertEqual([item["celula_valor"] for item in grupo["contas"]], ["B4", "B5"])

    def test_modelo_detectado_pode_alterar_e_excluir_contas_padrao(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha_dinamica(Path(temporario))
            modelo = detectar_modelo_saldo_contas(caminho, "OUT.26")["modelo"]
            modelo["nome"] = "Modelo editado"
            modelo["grupos"][0]["nome"] = "ABACATE ESPECIAL"
            modelo["grupos"][0]["contas"].pop(1)
            modelo = normalizar_modelo_saldo_contas(modelo)
            analise = analisar_planilha_saldo_contas(caminho, modelo=modelo)

        self.assertTrue(analise["modelo_valido"])
        self.assertEqual(analise["modelo_nome"], "Modelo editado")
        self.assertEqual(analise["grupos"][0]["nome"], "ABACATE ESPECIAL")
        self.assertEqual(analise["grupos"][0]["quantidade"], 1)

    def test_modelo_dinamico_preenche_copia_e_total_sem_alterar_formula(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha_dinamica(Path(temporario))
            original = caminho.read_bytes()
            modelo = detectar_modelo_saldo_contas(caminho, "OUT.26")["modelo"]
            saida = preencher_planilha_saldo_contas(
                caminho,
                {"11311001": Decimal("10.25"), "11311002": Decimal("20.75")},
                2,
                2028,
                grupos_selecionados=[modelo["grupos"][0]["id"]],
                modelo=modelo,
            )
            self.assertEqual(caminho.read_bytes(), original)
            workbook = load_workbook(saida["caminho_saida"], data_only=False)
            try:
                sheet = workbook["OUT.26"]
                self.assertEqual(sheet["B4"].value, 10.25)
                self.assertEqual(sheet["B5"].value, 20.75)
                self.assertEqual(sheet["A6"].value, 31)
                self.assertEqual(sheet["B6"].value, "=SUM(B4:B5)")
                self.assertEqual(sheet["A1"].value.date(), date(2028, 2, 1))
            finally:
                workbook.close()

    def test_modelo_bloqueia_destino_com_formula(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha_dinamica(Path(temporario))
            workbook = load_workbook(caminho)
            workbook["OUT.26"]["B4"] = "=1+1"
            workbook.save(caminho)
            workbook.close()
            modelo = detectar_modelo_saldo_contas(caminho, "OUT.26")["modelo"]
            primeira = modelo["grupos"][0]["contas"][0]
            primeira["ativo"] = True
            primeira["bloqueada"] = False
            analise = analisar_planilha_saldo_contas(caminho, modelo=modelo)

        self.assertFalse(analise["modelo_valido"])
        self.assertTrue(any("possui formula" in aviso for aviso in analise["avisos"]))

    def test_modelo_valido_com_41_contas(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(Path(temporario))
            resultado = analisar_planilha_saldo_contas(caminho)

        self.assertTrue(resultado["ok"])
        self.assertTrue(resultado["modelo_valido"])
        self.assertEqual(resultado["periodo"]["data"], "2026-09-01")
        self.assertEqual(resultado["resumo"]["contas_encontradas"], 41)
        self.assertEqual(resultado["resumo"]["saldos_preenchidos"], 41)
        self.assertEqual(resultado["resumo"]["conferencia"], "ok")

    def test_conta_diferente_bloqueia_modelo(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(
                Path(temporario),
                alterar_primeira_conta=True,
            )
            resultado = analisar_planilha_saldo_contas(caminho)

        self.assertTrue(resultado["ok"])
        self.assertFalse(resultado["modelo_valido"])
        self.assertTrue(any("Contas ausentes" in aviso for aviso in resultado["avisos"]))
        self.assertTrue(any("Contas diferentes" in aviso for aviso in resultado["avisos"]))

    def test_saldo_vazio_nao_e_convertido_em_zero(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(Path(temporario), vazia=True)
            resultado = analisar_planilha_saldo_contas(caminho)

        self.assertTrue(resultado["modelo_valido"])
        self.assertEqual(resultado["resumo"]["saldos_preenchidos"], 0)
        self.assertEqual(resultado["resumo"]["saldos_vazios"], 41)
        self.assertTrue(any("saldo(s) vazio(s)" in aviso for aviso in resultado["avisos"]))

    def test_grupos_e_amarelos_somente_na_coluna_da_conta(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(Path(temporario), vazia=True)
            wb = load_workbook(caminho)
            sh = wb["SET.26"]
            sh["A3"] = "BCO CTA MOVIMENTOS"
            sh["C3"] = "BCO CTA MOV. OUTRAS"
            sh["A22"] = "CENTRALIZADORA"
            sh["A4"] = 11311000
            sh["A6"] = 11311025
            sh["C4"] = 11311001
            sh["A23"] = 11311305
            sh["A24"] = 11311306
            sh["A6"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
            sh["A23"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
            sh["D20"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
            wb.save(caminho)
            wb.close()
            dados = analisar_planilha_saldo_contas(caminho)
            grupos = {g["id"]: g for g in dados["grupos"]}
            self.assertEqual([g["id"] for g in dados["grupos"]], [
                "bco_movimentos", "bco_outras", "centralizadora", "aplic_movimento",
            ])
            self.assertEqual(grupos["bco_movimentos"]["quantidade"], 1)
            self.assertEqual(grupos["bco_movimentos"]["individuais"], 1)
            self.assertEqual(grupos["centralizadora"]["individuais"], 1)
            self.assertEqual(grupos["aplic_movimento"]["quantidade"], 41)
            self.assertEqual(grupos["aplic_movimento"]["individuais"], 0)
            plano = planejar_consultas_saldo_contas(dados, [
                "bco_movimentos", "centralizadora", "aplic_movimento",
            ])
            self.assertEqual(len(plano["lotes"]), 4)  # 2 lotes + 2 solos
            self.assertEqual(len(plano["destinos"]), 44)
            self.assertNotIn("11311000", plano["destinos"])
            self.assertFalse(any(
                "11311000" in lote["contas"] for lote in plano["lotes"]
            ))
            self.assertEqual(plano["destinos"]["11311305"], "B23")
            self.assertEqual(plano["destinos"]["11370003"], "D22")

            # Consulta somente da Centralizadora: as outras 41 aplicacoes
            # continuam vazias, com formulas, estilos e periodo preservados.
            original = caminho.read_bytes()
            saida = preencher_planilha_saldo_contas(
                caminho,
                {"11311305": Decimal("1972978.12"), "11311306": Decimal("0")},
                9, 2026,
                grupos_selecionados=["centralizadora"],
            )
            self.assertEqual(caminho.read_bytes(), original)
            self.assertEqual(saida["contas_preenchidas"], 2)
            resultado = load_workbook(saida["caminho_saida"], data_only=False)
            try:
                ws = resultado["SET.26"]
                self.assertEqual(ws["B23"].value, 1972978.12)
                self.assertEqual(ws["B24"].value, 0)
                self.assertIsNone(ws["D20"].value)
                self.assertEqual(ws["D61"].value, "=SUM(D20:D60)")
                self.assertIsNone(ws["C61"].value)
            finally:
                resultado.close()

    def test_grupo_desconhecido_e_consulta_incompleta_bloqueiam_escrita(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(Path(temporario), vazia=True)
            analise = analisar_planilha_saldo_contas(caminho)
            with self.assertRaisesRegex(ValueError, "indisponivel"):
                planejar_consultas_saldo_contas(analise, ["fundo_fixo"])
            with self.assertRaisesRegex(ValueError, "Selecione"):
                planejar_consultas_saldo_contas(analise, [])
            with self.assertRaisesRegex(ValueError, "apenas 1 de 41"):
                preencher_planilha_saldo_contas(
                    caminho, {"11370000": 0}, 9, 2026,
                    grupos_selecionados=["aplic_movimento"],
                )
            self.assertEqual(len(list(Path(temporario).glob("*.xlsx"))), 1)

    def test_exige_aba_set_26(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = Path(temporario) / "sem-modelo.xlsx"
            workbook = Workbook()
            workbook.active.title = "OUT.26"
            workbook.save(caminho)
            workbook.close()
            resultado = analisar_planilha_saldo_contas(caminho)

        self.assertFalse(resultado["ok"])
        self.assertIn("SET.26", resultado["msg"])

    def test_bloqueia_planilha_incompleta_sem_gerar_arquivo(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(Path(temporario), vazia=True)
            original = caminho.read_bytes()
            with self.assertRaisesRegex(ValueError, "apenas 1 de 41 contas"):
                preencher_planilha_saldo_contas(
                    caminho, {"11370000": Decimal("0")}, 9, 2026
                )
            self.assertEqual(caminho.read_bytes(), original)
            self.assertEqual(len(list(Path(temporario).glob("*.xlsx"))), 1)

    def test_preenche_copia_preservando_original_formulas_e_estilo(self):
        with tempfile.TemporaryDirectory() as temporario:
            caminho = self._criar_planilha(Path(temporario), vazia=True)
            original = caminho.read_bytes()
            saldos = {conta: Decimal("0") for conta in CONTAS_APLICACAO_MODELO}
            saldos["11370000"] = Decimal("3394757.09")
            saldos["11370002"] = Decimal("-25.01")
            resultado = preencher_planilha_saldo_contas(caminho, saldos, 2, 2028)
            saida = Path(resultado["caminho_saida"])

            self.assertTrue(saida.exists())
            self.assertEqual(caminho.read_bytes(), original)
            self.assertEqual(resultado["contas_retornadas"], 41)
            self.assertEqual(resultado["total"], 3394732.08)

            workbook = load_workbook(saida, data_only=False)

            try:
                sheet = workbook["SET.26"]
                self.assertEqual(sheet["D20"].value, 3394757.09)
                self.assertEqual(sheet["D21"].value, -25.01)
                self.assertEqual(sheet["D22"].value, 0)
                self.assertEqual(sheet["C61"].value, 3394732.08)
                self.assertEqual(sheet["D61"].value, "=SUM(D20:D60)")
                self.assertEqual(sheet["B59"].value, "=D59+D60")
                self.assertEqual(sheet["J7"].value, "=J3+J4+J5+J6")
                self.assertEqual(sheet["C1"].value.date(), date(2028, 2, 1))
                self.assertEqual(sheet["D20"].fill.fgColor.rgb, "00FFF200")
                self.assertEqual(
                    sheet["D20"].number_format,
                    '#,##0.00;[Red]-#,##0.00;"-"',
                )
            finally:
                workbook.close()

if __name__ == "__main__":
    unittest.main()
