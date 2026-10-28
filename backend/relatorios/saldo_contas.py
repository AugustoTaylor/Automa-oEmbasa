from __future__ import annotations

import math
import os
import re
import unicodedata
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from uuid import uuid4


ABA_MODELO_SALDO_CONTAS = "SET.26"
LINHA_INICIAL_CONTAS = 20
LINHA_FINAL_CONTAS = 60
TOTAL_CONTAS_MODELO = LINHA_FINAL_CONTAS - LINHA_INICIAL_CONTAS + 1
LIMITE_TAMANHO_PLANILHA_BYTES = 50 * 1024 * 1024

# Intervalos fixos do modelo enviado: conta em A/C e saldo a direita.
# A cor amarela so vale na CELULA DO NUMERO DA CONTA. Uma coluna de saldos
# amarela (por exemplo D20:D58) nao torna as 41 aplicacoes individuais.
GRUPOS_CONTAS = (
    ("bco_movimentos", "BCO CTA MOVIMENTOS", 1, 2, 4, 17, "A3"),
    ("bco_outras", "BCO CTA MOV. OUTRAS", 3, 4, 4, 17, "C3"),
    ("centralizadora", "CENTRALIZADORA", 1, 2, 23, 28, "A22"),
    ("aplic_movimento", "APLIC. MOVIMENTO", 3, 4, 20, 60, "C19"),
)
GRUPOS_IDS = frozenset(grupo[0] for grupo in GRUPOS_CONTAS)

# A conta 11311000 possui volume historico excessivo no SAP e causa timeout.
# Ela permanece visivel no modelo, mas nao participa da consulta nem da escrita.
CONTAS_IGNORADAS_SAP = frozenset({"11311000"})


CONTAS_APLICACAO_MODELO = (
    "11370000",
    "11370002",
    "11370003",
    "11370004",
    "11370005",
    "11370006",
    "11370007",
    "11370008",
    "11370009",
    "11370010",
    "11370011",
    "11370012",
    "11370013",
    "11371001",
    "11371003",
    "11371065",
    "11371067",
    "11371068",
    "11371069",
    "11371070",
    "11371071",
    "11371072",
    "11371073",
    "11371074",
    "11371075",
    "11371076",
    "11371077",
    "11371078",
    "11371079",
    "11371081",
    "11371082",
    "11371083",
    "11371301",
    "11371353",
    "11371500",
    "11371601",
    "11371602",
    "11374005",
    "11374006",
    "11374007",
    "11374008",
)

_ABAS_MENSAIS_RE = re.compile(
    r"^(JAN|FEV|MAR|ABR|MAI|JUN|JUL|AGO|SET|OUT|NOV|DEZ)\.(\d{2})$",
    re.IGNORECASE,
)
_FORMULAS_ESPERADAS = {
    "D61": "=SUM(D20:D60)",
    "B59": "=D59+D60",
    "J5": "=D61-J6",
    "J6": "=B59",
    "J7": "=J3+J4+J5+J6",
}
_CENTAVOS = Decimal("0.01")
_CELULA_A1_RE = re.compile(r"^[A-Z]{1,3}[1-9]\d*$")
_CABECALHOS_GENERICOS = frozenset({
    "INDIVIDUAIS",
    "INDIVIDUAL",
    "VALOR",
    "VALORES",
    "SALDO",
    "SALDOS",
    "CONTA",
    "CONTAS",
    "CONTA DO RAZAO",
    "R$",
})
_MAX_LINHAS_DETECCAO = 5000
_MAX_COLUNAS_DETECCAO = 200


def _normalizar_texto(valor) -> str:
    return " ".join(str(valor or "").strip().upper().split())


def _normalizar_formula(valor) -> str:
    return re.sub(r"[\s$]", "", str(valor or "")).upper()


def _normalizar_conta(valor) -> str:
    if isinstance(valor, bool) or valor is None:
        return ""

    if isinstance(valor, int):
        return str(valor)

    if isinstance(valor, float):
        if not math.isfinite(valor) or not valor.is_integer():
            return ""
        return str(int(valor))

    texto = str(valor).strip()

    if re.fullmatch(r"\d+(?:\.0+)?", texto):
        return texto.split(".", 1)[0]

    return re.sub(r"\D", "", texto)


def _conta_exata_celula(valor) -> str:
    """Aceita somente uma celula cujo conteudo inteiro representa a conta."""
    if isinstance(valor, bool) or valor is None:
        return ""
    if isinstance(valor, int):
        conta = str(valor)
    elif isinstance(valor, float):
        conta = str(int(valor)) if math.isfinite(valor) and valor.is_integer() else ""
    else:
        texto = str(valor).strip()
        correspondencia = re.fullmatch(r"(\d{8})(?:\.0+)?", texto)
        conta = correspondencia.group(1) if correspondencia else ""
    return conta if re.fullmatch(r"\d{8}", conta) else ""


def _slug_modelo(valor, fallback="grupo") -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    texto = re.sub(r"[^a-zA-Z0-9]+", "_", texto).strip("_").lower()
    return texto[:48] or fallback


def _celula_a1(valor, *, obrigatoria=True) -> str:
    texto = str(valor or "").strip().upper().replace("$", "")
    if not texto and not obrigatoria:
        return ""
    if not _CELULA_A1_RE.fullmatch(texto):
        raise ValueError(f"Endereco de celula invalido: {valor or '--'}.")
    return texto


def normalizar_modelo_saldo_contas(modelo) -> dict:
    """Valida e normaliza um mapeamento editavel de grupos e contas."""
    if not isinstance(modelo, dict):
        raise ValueError("O modelo de contas enviado e invalido.")

    resultado = deepcopy(modelo)
    resultado["id"] = str(resultado.get("id") or "").strip()
    resultado["nome"] = str(resultado.get("nome") or "").strip()[:90]
    resultado["aba"] = str(resultado.get("aba") or "").strip()[:31]
    resultado["celula_periodo"] = _celula_a1(
        resultado.get("celula_periodo"), obrigatoria=False
    )

    if not resultado["nome"]:
        raise ValueError("Informe um nome para o modelo de planilha.")
    if not resultado["aba"]:
        raise ValueError("Selecione a aba usada pelo modelo.")

    grupos_brutos = resultado.get("grupos")
    if not isinstance(grupos_brutos, list) or not grupos_brutos:
        raise ValueError("O modelo precisa possuir pelo menos um grupo.")

    grupos = []
    ids_grupos = set()
    contas_vistas = set()
    destinos_ativos = set()
    total_ativas = 0

    for indice_grupo, grupo_bruto in enumerate(grupos_brutos, start=1):
        if not isinstance(grupo_bruto, dict):
            raise ValueError(f"O grupo {indice_grupo} e invalido.")
        nome = str(grupo_bruto.get("nome") or "").strip()[:90]
        if not nome:
            raise ValueError(f"Informe o nome do grupo {indice_grupo}.")
        base_id = _slug_modelo(grupo_bruto.get("id") or nome, f"grupo_{indice_grupo}")
        grupo_id = base_id
        sufixo = 2
        while grupo_id in ids_grupos:
            grupo_id = f"{base_id}_{sufixo}"
            sufixo += 1
        ids_grupos.add(grupo_id)

        grupo = {
            "id": grupo_id,
            "nome": nome,
            "ativo": bool(grupo_bruto.get("ativo", True)),
            "celula_cabecalho": _celula_a1(
                grupo_bruto.get("celula_cabecalho"), obrigatoria=False
            ),
            "celula_total": _celula_a1(
                grupo_bruto.get("celula_total"), obrigatoria=False
            ),
            "contas": [],
        }
        contas_brutas = grupo_bruto.get("contas")
        if not isinstance(contas_brutas, list):
            raise ValueError(f"A lista de contas de {nome} e invalida.")

        for indice_conta, item_bruto in enumerate(contas_brutas, start=1):
            if not isinstance(item_bruto, dict):
                raise ValueError(f"A conta {indice_conta} de {nome} e invalida.")
            conta = _normalizar_conta(item_bruto.get("conta"))
            if not re.fullmatch(r"\d{8}", conta):
                raise ValueError(
                    f"A conta {indice_conta} de {nome} deve possuir exatamente 8 digitos."
                )
            if conta in contas_vistas:
                raise ValueError(f"Conta duplicada no modelo: {conta}.")
            contas_vistas.add(conta)

            celula_conta = _celula_a1(item_bruto.get("celula_conta"))
            celula_valor = _celula_a1(
                item_bruto.get("celula_valor") or item_bruto.get("celula")
            )
            if celula_conta == celula_valor:
                raise ValueError(f"A conta {conta} nao pode receber o saldo na propria celula.")

            modo = str(item_bruto.get("modo") or "").strip().lower()
            if not modo:
                modo = "individual" if item_bruto.get("individual") else "lote"
            if modo not in {"lote", "individual"}:
                raise ValueError(f"Modo de consulta invalido para a conta {conta}.")

            bloqueada = conta in CONTAS_IGNORADAS_SAP
            ativo = bool(item_bruto.get("ativo", True)) and not bloqueada
            ativo_efetivo = ativo and grupo["ativo"]
            if ativo_efetivo and celula_valor in destinos_ativos:
                raise ValueError(f"Mais de uma conta usa a celula de valor {celula_valor}.")
            if ativo_efetivo:
                destinos_ativos.add(celula_valor)
                total_ativas += 1

            grupo["contas"].append({
                "id": str(item_bruto.get("id") or f"{grupo_id}_{conta}_{celula_conta}")[:140],
                "conta": conta,
                "celula_conta": celula_conta,
                "celula_valor": celula_valor,
                "celula": celula_valor,
                "modo": modo,
                "individual": modo == "individual",
                "ativo": ativo,
                "bloqueada": bloqueada,
                "motivo_bloqueio": (
                    "Conta desativada por causar timeout no SAP."
                    if bloqueada else ""
                ),
            })

        grupos.append(grupo)

    if total_ativas == 0:
        raise ValueError("Ative pelo menos uma conta valida no modelo.")

    resultado["grupos"] = grupos
    resultado["versao"] = 1
    return resultado


def _decimal(valor) -> Decimal | None:
    if isinstance(valor, bool) or valor is None or valor == "":
        return None

    try:
        numero = Decimal(str(valor))
    except (InvalidOperation, ValueError, TypeError):
        return None

    if not numero.is_finite():
        return None

    return numero


def _decimal_centavos(valor) -> Decimal | None:
    numero = _decimal(valor)
    return numero.quantize(_CENTAVOS, rounding=ROUND_HALF_UP) if numero is not None else None


def _numero_front(valor: Decimal | None):
    return float(valor) if valor is not None else None


def _valor_formula_constante(valor) -> Decimal | None:
    correspondencia = re.fullmatch(r"=\s*(-?\d+(?:[.,]\d+)?)\s*", str(valor or ""))

    if not correspondencia:
        return None

    return _decimal_centavos(correspondencia.group(1).replace(",", "."))


def _periodo_planilha(valor):
    if isinstance(valor, datetime):
        referencia = valor.date()
    elif isinstance(valor, date):
        referencia = valor
    else:
        referencia = None

    if referencia is None:
        return {"data": "", "mes": "", "ano": ""}

    return {
        "data": referencia.isoformat(),
        "mes": f"{referencia.month:02d}",
        "ano": str(referencia.year),
    }


def _validar_arquivo(caminho) -> Path:
    texto = str(caminho or "").strip()

    if not texto:
        raise ValueError("Selecione uma planilha Excel para continuar.")

    arquivo = Path(texto).expanduser()

    if not arquivo.exists() or not arquivo.is_file():
        raise ValueError("A planilha selecionada nao foi encontrada.")

    if arquivo.suffix.lower() != ".xlsx":
        raise ValueError("Selecione uma planilha no formato .xlsx.")

    if arquivo.stat().st_size > LIMITE_TAMANHO_PLANILHA_BYTES:
        raise ValueError("A planilha ultrapassa o limite de 50 MB.")

    return arquivo.resolve()


def _carregar_planilha(caminho: Path, *, data_only: bool):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "O componente de leitura de planilhas nao esta instalado nesta versao do aplicativo."
        ) from exc

    return load_workbook(
        filename=caminho,
        read_only=True,
        data_only=data_only,
        keep_links=False,
    )


def _carregar_planilha_edicao(caminho: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "O componente de escrita de planilhas nao esta instalado nesta versao do aplicativo."
        ) from exc

    return load_workbook(
        filename=caminho,
        data_only=False,
        keep_links=True,
    )


def _proximo_caminho_saida(arquivo: Path) -> Path:
    candidato = arquivo.with_name(f"{arquivo.stem} - PREENCHIDO{arquivo.suffix}")

    if not candidato.exists():
        return candidato

    indice = 2

    while True:
        candidato = arquivo.with_name(
            f"{arquivo.stem} - PREENCHIDO ({indice}){arquivo.suffix}"
        )

        if not candidato.exists():
            return candidato

        indice += 1



def _conta_amarela(celula) -> bool:
    """Somente destaque amarelo do identificador contabil determina consulta solo."""
    try:
        preenchimento = celula.fill
        if preenchimento.patternType != "solid":
            return False
        cor = preenchimento.fgColor
        if cor.type == "rgb" and isinstance(cor.rgb, str):
            return cor.rgb[-6:].upper() in {"FFFF00", "FFF200", "FFEB00"}
        if cor.type == "indexed":
            return cor.indexed in {5, 6, 13}
    except (AttributeError, TypeError):
        return False
    return False


def _texto_cabecalho(valor) -> str:
    if not isinstance(valor, str) or str(valor).lstrip().startswith("="):
        return ""
    texto = _normalizar_texto(valor)
    if not texto or texto in _CABECALHOS_GENERICOS:
        return ""
    if re.fullmatch(r"[\d\s.,:/-]+", texto):
        return ""
    return " ".join(str(valor).strip().split())[:90]


def _cabecalho_mais_proximo(sheet, primeira_linha: int, coluna: int):
    inicio = max(1, primeira_linha - 15)
    for linha in range(primeira_linha - 1, inicio - 1, -1):
        texto = _texto_cabecalho(sheet.cell(linha, coluna).value)
        if texto:
            return texto, sheet.cell(linha, coluna).coordinate
    for deslocamento in (-1, 1):
        coluna_vizinha = coluna + deslocamento
        if coluna_vizinha < 1:
            continue
        for linha in range(primeira_linha - 1, inicio - 1, -1):
            texto = _texto_cabecalho(sheet.cell(linha, coluna_vizinha).value)
            if texto:
                return texto, sheet.cell(linha, coluna_vizinha).coordinate
    return f"GRUPO {sheet.cell(primeira_linha, coluna).column_letter}", ""


def _escolher_aba_modelo(workbook, aba_modelo=None) -> str:
    solicitada = str(aba_modelo or "").strip()
    if solicitada:
        if solicitada not in workbook.sheetnames:
            raise ValueError(f'A aba "{solicitada}" nao foi encontrada na planilha.')
        return solicitada
    if ABA_MODELO_SALDO_CONTAS in workbook.sheetnames:
        return ABA_MODELO_SALDO_CONTAS
    mensais = [nome for nome in workbook.sheetnames if _ABAS_MENSAIS_RE.fullmatch(nome)]
    return mensais[-1] if mensais else workbook.sheetnames[0]


def detectar_modelo_saldo_contas(caminho, aba_modelo=None) -> dict:
    """Varre uma aba e propoe grupos/contas para confirmacao do usuario."""
    arquivo = _validar_arquivo(caminho)
    workbook = None
    try:
        workbook = _carregar_planilha_edicao(arquivo)
        aba = _escolher_aba_modelo(workbook, aba_modelo)
        sheet = workbook[aba]
        max_row = min(int(sheet.max_row or 1), _MAX_LINHAS_DETECCAO)
        max_column = min(int(sheet.max_column or 1), _MAX_COLUNAS_DETECCAO)
        por_coluna = {}

        for linha in range(1, max_row + 1):
            for coluna in range(1, max_column + 1):
                celula = sheet.cell(linha, coluna)
                conta = _conta_exata_celula(celula.value)
                if conta:
                    por_coluna.setdefault(coluna, []).append((linha, celula, conta))

        grupos = []
        ids_usados = set()
        avisos = []
        for coluna, registros in sorted(por_coluna.items()):
            segmentos = []
            atual = []
            linha_anterior = None
            for registro in registros:
                linha = registro[0]
                cabecalho_entre = False
                if linha_anterior is not None:
                    for intermediaria in range(linha_anterior + 1, linha):
                        if _texto_cabecalho(sheet.cell(intermediaria, coluna).value):
                            cabecalho_entre = True
                            break
                if atual and (linha - linha_anterior > 3 or cabecalho_entre):
                    segmentos.append(atual)
                    atual = []
                atual.append(registro)
                linha_anterior = linha
            if atual:
                segmentos.append(atual)

            for segmento in segmentos:
                # Blocos verticais com pelo menos duas contas sao deteccoes
                # confiaveis. Contas isoladas e intervalos permanecem disponiveis
                # para cadastro manual no editor, evitando falsos positivos.
                if len(segmento) < 2:
                    continue
                primeira_linha = segmento[0][0]
                ultima_linha = segmento[-1][0]
                nome, celula_cabecalho = _cabecalho_mais_proximo(
                    sheet, primeira_linha, coluna
                )
                base_id = _slug_modelo(nome, f"grupo_{len(grupos) + 1}")
                grupo_id = base_id
                sufixo = 2
                while grupo_id in ids_usados:
                    grupo_id = f"{base_id}_{sufixo}"
                    sufixo += 1
                ids_usados.add(grupo_id)
                contas = []
                for linha, celula, conta in segmento:
                    destino = sheet.cell(linha, coluna + 1)
                    formula_destino = (
                        isinstance(destino.value, str)
                        and destino.value.lstrip().startswith("=")
                    )
                    bloqueada = conta in CONTAS_IGNORADAS_SAP or formula_destino
                    motivo = ""
                    if conta in CONTAS_IGNORADAS_SAP:
                        motivo = "Conta desativada por causar timeout no SAP."
                        avisos.append(
                            f"A conta {conta} foi encontrada em {celula.coordinate} e permaneceu desativada."
                        )
                    elif formula_destino:
                        motivo = f"A celula {destino.coordinate} possui formula."
                        avisos.append(
                            f"A conta {conta} foi desativada porque {destino.coordinate} possui formula."
                        )
                    contas.append({
                        "id": f"{grupo_id}_{conta}_{celula.coordinate}",
                        "conta": conta,
                        "celula_conta": celula.coordinate,
                        "celula_valor": destino.coordinate,
                        "celula": destino.coordinate,
                        "modo": "individual" if _conta_amarela(celula) else "lote",
                        "individual": _conta_amarela(celula),
                        "ativo": not bloqueada,
                        "bloqueada": bloqueada,
                        "motivo_bloqueio": motivo,
                    })

                linha_total = ultima_linha + 1
                celula_total = ""
                if linha_total <= max_row + 1:
                    candidato_total = sheet.cell(linha_total, coluna)
                    candidato_formula = sheet.cell(linha_total, coluna + 1)
                    total_ja_e_formula = (
                        isinstance(candidato_total.value, str)
                        and candidato_total.value.lstrip().startswith("=")
                    )
                    if not total_ja_e_formula and (
                        candidato_total.value not in (None, "")
                        or (
                            isinstance(candidato_formula.value, str)
                            and candidato_formula.value.lstrip().startswith("=")
                        )
                    ):
                        celula_total = candidato_total.coordinate

                grupos.append({
                    "id": grupo_id,
                    "nome": nome,
                    "ativo": True,
                    "celula_cabecalho": celula_cabecalho,
                    "celula_total": celula_total,
                    "contas": contas,
                })

        if not grupos:
            raise ValueError(
                "Nenhuma sequencia de contas SAP com 8 digitos foi encontrada na aba selecionada."
            )

        celula_periodo = ""
        if isinstance(sheet["C1"].value, (date, datetime)):
            celula_periodo = "C1"
        else:
            for linha in range(1, min(max_row, 10) + 1):
                for coluna in range(1, min(max_column, 10) + 1):
                    if isinstance(sheet.cell(linha, coluna).value, (date, datetime)):
                        celula_periodo = sheet.cell(linha, coluna).coordinate
                        break
                if celula_periodo:
                    break

        modelo = normalizar_modelo_saldo_contas({
            "id": "",
            "nome": f"Saldo de contas - {aba}",
            "aba": aba,
            "celula_periodo": celula_periodo,
            "grupos": grupos,
        })
        total_contas = sum(len(grupo["contas"]) for grupo in modelo["grupos"])
        total_ativas = sum(
            item["ativo"] for grupo in modelo["grupos"] for item in grupo["contas"]
        )
        return {
            "ok": True,
            "msg": (
                f"{total_contas} conta(s) identificada(s) em "
                f"{len(modelo['grupos'])} grupo(s). Confirme o modelo antes de consultar."
            ),
            "caminho": str(arquivo),
            "arquivo": arquivo.name,
            "abas": list(workbook.sheetnames),
            "aba_modelo": aba,
            "modelo": modelo,
            "modelo_valido": total_ativas > 0,
            "requer_confirmacao": True,
            "avisos": list(dict.fromkeys(avisos)),
            "resumo": {
                "contas_encontradas": total_contas,
                "contas_ativas": total_ativas,
                "grupos_encontrados": len(modelo["grupos"]),
            },
        }
    finally:
        if workbook is not None:
            workbook.close()


def ler_grupos_contas_modelo(sheet) -> list[dict]:
    """Identifica contas e destinos no modelo, sem incluir subtotais/intervalos."""
    grupos = []
    contas_ja_vistas = set()
    for chave, nome, coluna_conta, coluna_valor, inicio, fim, cabecalho in GRUPOS_CONTAS:
        titulo = _normalizar_texto(sheet[cabecalho].value)
        if titulo != _normalizar_texto(nome):
            # Um modelo incompleto nao pode causar lancamentos em locais errados.
            if chave == "aplic_movimento":
                raise ValueError("Cabecalho APLIC. MOVIMENTO nao encontrado em C19.")
            continue
        itens = []
        for linha in range(inicio, fim + 1):
            celula_conta = sheet.cell(linha, coluna_conta)
            conta = _normalizar_conta(celula_conta.value)
            if not re.fullmatch(r"\d{8}", conta):
                continue
            if conta in CONTAS_IGNORADAS_SAP:
                continue
            if conta in contas_ja_vistas:
                raise ValueError(f"Conta duplicada em mais de um grupo: {conta}.")
            contas_ja_vistas.add(conta)
            itens.append({
                "conta": conta,
                "celula": f"{chr(64 + coluna_valor)}{linha}",
                "individual": _conta_amarela(celula_conta),
            })
        if itens:
            grupos.append({
                "id": chave,
                "nome": nome,
                "quantidade": len(itens),
                "individuais": sum(item["individual"] for item in itens),
                "em_lote": sum(not item["individual"] for item in itens),
                "contas": itens,
            })
    return grupos


def _validar_grupos_selecionados(grupos_disponiveis, selecionados):
    if selecionados is None:
        selecionados = ("aplic_movimento",)  # compatibilidade com fluxo anterior
    if not isinstance(selecionados, (list, tuple, set, frozenset)):
        raise ValueError("A selecao de grupos deve ser uma lista.")
    ids = list(selecionados)
    if not ids or any(not isinstance(chave, str) for chave in ids):
        raise ValueError("Selecione pelo menos um grupo de contas.")
    if len(set(ids)) != len(ids):
        raise ValueError("Ha grupos repetidos na selecao.")
    disponiveis = {item["id"]: item for item in grupos_disponiveis}
    desconhecidos = set(ids) - disponiveis.keys()
    if desconhecidos:
        raise ValueError(f"Grupo(s) indisponivel(is) nesta planilha: {', '.join(sorted(desconhecidos))}.")
    return [disponiveis[chave] for chave in ids]


def planejar_consultas_saldo_contas(analise, selecionados=None) -> dict:
    """Agrupa contas normais e separa cada conta destacada em amarelo."""
    grupos = _validar_grupos_selecionados(analise.get("grupos") or [], selecionados)
    lotes = []
    destinos = {}
    for grupo in grupos:
        contas_lote = []
        for item in grupo["contas"]:
            destinos[item["conta"]] = item["celula"]
            if not item["individual"]:
                contas_lote.append(item["conta"])
        if contas_lote:
            lotes.append({"grupo": grupo["nome"], "contas": contas_lote, "individual": False})
        for item in grupo["contas"]:
            if item["individual"]:
                lotes.append({"grupo": grupo["nome"], "contas": [item["conta"]], "individual": True})
    return {"grupos": [g["id"] for g in grupos], "lotes": lotes, "destinos": destinos}


def preencher_planilha_saldo_contas(
    caminho,
    saldos_por_conta,
    mes,
    ano,
    aba_modelo: str = ABA_MODELO_SALDO_CONTAS,
    grupos_selecionados=None,
    modelo=None,
) -> dict:
    arquivo = _validar_arquivo(caminho)
    configuracao = normalizar_modelo_saldo_contas(modelo) if modelo is not None else None
    analise_origem = analisar_planilha_saldo_contas(
        arquivo,
        aba_modelo=aba_modelo,
        modelo=configuracao,
    )

    if not analise_origem.get("ok") or not analise_origem.get("modelo_valido"):
        raise ValueError(
            "A planilha selecionada nao corresponde ao modelo de contas confirmado."
        )

    try:
        mes_numero = int(mes)
        ano_numero = int(ano)
    except (TypeError, ValueError) as exc:
        raise ValueError("Informe mes e ano validos para preencher a planilha.") from exc

    if mes_numero < 1 or mes_numero > 12:
        raise ValueError("O mes do relatorio deve estar entre 1 e 12.")

    if ano_numero < 2020 or ano_numero > 2099:
        raise ValueError("O ano do relatorio deve estar entre 2020 e 2099.")

    plano = planejar_consultas_saldo_contas(analise_origem, grupos_selecionados)
    destinos = plano["destinos"]
    saldos_normalizados = {}
    for conta_bruta, saldo_bruto in dict(saldos_por_conta or {}).items():
        conta = _normalizar_conta(conta_bruta)
        if conta not in destinos:
            continue
        saldo = _decimal_centavos(saldo_bruto)
        if saldo is None:
            raise ValueError(f"Saldo invalido retornado para a conta {conta}.")
        saldos_normalizados[conta] = saldo

    contas_ausentes = sorted(destinos.keys() - saldos_normalizados.keys())
    if contas_ausentes:
        raise ValueError(
            f"O SAP retornou apenas {len(saldos_normalizados)} de "
            f"{len(destinos)} contas selecionadas. Planilha nao modificada. "
            f"Faltantes: {', '.join(contas_ausentes)}."
        )

    workbook = None
    caminho_saida = _proximo_caminho_saida(arquivo)
    caminho_temporario = caminho_saida.with_name(
        f".{caminho_saida.stem}.{uuid4().hex}.tmp.xlsx"
    )

    try:
        workbook = _carregar_planilha_edicao(arquivo)
        aba_destino = configuracao["aba"] if configuracao is not None else aba_modelo
        sheet = workbook[aba_destino]
        total = Decimal("0")

        for conta, referencia in destinos.items():
            saldo = saldos_normalizados[conta]
            sheet[referencia].value = float(saldo)
            total += saldo

        total = total.quantize(_CENTAVOS, rounding=ROUND_HALF_UP)
        if configuracao is not None:
            grupos_selecionados_set = set(plano["grupos"])
            for grupo in analise_origem["grupos"]:
                if grupo["id"] not in grupos_selecionados_set or not grupo.get("celula_total"):
                    continue
                total_grupo = sum(
                    (saldos_normalizados[item["conta"]] for item in grupo["contas"]),
                    Decimal("0"),
                ).quantize(_CENTAVOS, rounding=ROUND_HALF_UP)
                sheet[grupo["celula_total"]] = float(total_grupo)
        elif "aplic_movimento" in plano["grupos"]:
            # A celula C61 sempre foi o subtotal independente das 41 aplicacoes.
            # Nao alteramos essa celula se o usuario consultou apenas outros grupos.
            total_aplic = sum(
                (saldos_normalizados[item["conta"]]
                 for grupo in analise_origem["grupos"]
                 if grupo["id"] == "aplic_movimento"
                 for item in grupo["contas"]),
                Decimal("0"),
            )
            sheet["C61"] = float(total_aplic.quantize(_CENTAVOS))
        if configuracao is not None:
            if configuracao["celula_periodo"]:
                sheet[configuracao["celula_periodo"]] = date(ano_numero, mes_numero, 1)
        else:
            sheet["C1"] = date(ano_numero, mes_numero, 1)

        calculo = getattr(workbook, "calculation", None)

        if calculo is not None:
            calculo.fullCalcOnLoad = True
            calculo.forceFullCalc = True
            calculo.calcMode = "auto"

        workbook.save(caminho_temporario)
        workbook.close()
        workbook = None
        os.replace(caminho_temporario, caminho_saida)
    finally:
        if workbook is not None:
            workbook.close()

        if caminho_temporario.exists():
            try:
                caminho_temporario.unlink()
            except OSError:
                pass

    return {
        "ok": True,
        "msg": "Planilha preenchida e salva com sucesso.",
        "caminho_origem": str(arquivo),
        "caminho_saida": str(caminho_saida),
        "arquivo_saida": caminho_saida.name,
        "aba_modelo": configuracao["aba"] if configuracao is not None else aba_modelo,
        "modelo_id": configuracao.get("id") if configuracao is not None else "",
        "modelo_nome": configuracao.get("nome") if configuracao is not None else "Modelo SET.26",
        "periodo": {
            "mes": f"{mes_numero:02d}",
            "ano": str(ano_numero),
        },
        "grupos_selecionados": plano["grupos"],
        "contas_preenchidas": len(destinos),
        "contas_retornadas": len(saldos_normalizados),
        "contas_zeradas": sum(valor == 0 for valor in saldos_normalizados.values()),
        "total": _numero_front(total),
    }


def _criar_verificacao(codigo: str, titulo: str, ok: bool, detalhe: str) -> dict:
    return {
        "codigo": codigo,
        "titulo": titulo,
        "ok": bool(ok),
        "detalhe": detalhe,
    }


def _analisar_planilha_com_modelo(caminho, modelo) -> dict:
    arquivo = _validar_arquivo(caminho)
    configuracao = normalizar_modelo_saldo_contas(modelo)
    workbook_formulas = None
    workbook_valores = None
    try:
        workbook_formulas = _carregar_planilha(arquivo, data_only=False)
        workbook_valores = _carregar_planilha(arquivo, data_only=True)
        abas = list(workbook_formulas.sheetnames)
        aba = configuracao["aba"]
        if aba not in workbook_formulas.sheetnames:
            erro_aba = f'A aba configurada "{aba}" nao foi encontrada na planilha.'
            return {
                "ok": False,
                "msg": erro_aba,
                "caminho": str(arquivo),
                "arquivo": arquivo.name,
                "abas": abas,
                "modelo": configuracao,
                "modelo_valido": False,
                "erros_modelo": [erro_aba],
            }

        sheet_formulas = workbook_formulas[aba]
        sheet_valores = workbook_valores[aba]
        erros = []
        avisos = []
        grupos_disponiveis = []
        contas_resumo = []
        saldos = []
        total_desativadas = 0

        for grupo in configuracao["grupos"]:
            itens_ativos = []
            for item in grupo["contas"]:
                if not item["ativo"] or not grupo["ativo"]:
                    total_desativadas += 1
                    continue
                conta_encontrada = _normalizar_conta(
                    sheet_formulas[item["celula_conta"]].value
                )
                if conta_encontrada != item["conta"]:
                    erros.append(
                        f"{item['celula_conta']}: esperada a conta {item['conta']}, "
                        f"mas foi encontrado {conta_encontrada or 'vazio'}."
                    )
                    continue
                destino = sheet_formulas[item["celula_valor"]]
                if isinstance(destino.value, str) and destino.value.lstrip().startswith("="):
                    erros.append(
                        f"{item['celula_valor']}: a celula de valor da conta "
                        f"{item['conta']} possui formula e nao sera sobrescrita."
                    )
                    continue
                saldo = _decimal_centavos(sheet_valores[item["celula_valor"]].value)
                saldos.append(saldo)
                contas_resumo.append({
                    "conta": item["conta"],
                    "celula_conta": item["celula_conta"],
                    "celula": item["celula_valor"],
                    "saldo_atual": _numero_front(saldo),
                    "status": "vazio" if saldo is None else ("zero" if saldo == 0 else "preenchido"),
                    "grupo": grupo["nome"],
                    "individual": item["individual"],
                })
                itens_ativos.append({
                    "conta": item["conta"],
                    "celula_conta": item["celula_conta"],
                    "celula": item["celula_valor"],
                    "celula_valor": item["celula_valor"],
                    "individual": item["individual"],
                    "modo": item["modo"],
                    "ativo": True,
                })

            if grupo["ativo"] and itens_ativos:
                if grupo["celula_total"]:
                    total_atual = sheet_formulas[grupo["celula_total"]].value
                    if isinstance(total_atual, str) and total_atual.lstrip().startswith("="):
                        erros.append(
                            f"{grupo['celula_total']}: o total manual de {grupo['nome']} "
                            "possui formula e nao sera sobrescrito."
                        )
                grupos_disponiveis.append({
                    "id": grupo["id"],
                    "nome": grupo["nome"],
                    "quantidade": len(itens_ativos),
                    "individuais": sum(item["individual"] for item in itens_ativos),
                    "em_lote": sum(not item["individual"] for item in itens_ativos),
                    "celula_total": grupo["celula_total"],
                    "contas": itens_ativos,
                })

        if not grupos_disponiveis:
            erros.append("O modelo nao possui contas ativas compativeis com esta aba.")

        if configuracao["celula_periodo"]:
            valor_periodo = sheet_formulas[configuracao["celula_periodo"]].value
            if isinstance(valor_periodo, str) and valor_periodo.lstrip().startswith("="):
                erros.append(
                    f"{configuracao['celula_periodo']}: a celula do periodo possui formula."
                )
            periodo = _periodo_planilha(valor_periodo)
        else:
            periodo = {"data": "", "mes": "", "ano": ""}

        preenchidos = sum(saldo is not None for saldo in saldos)
        zerados = sum(saldo == 0 for saldo in saldos)
        total_atual = sum((saldo or Decimal("0") for saldo in saldos), Decimal("0"))
        total_atual = total_atual.quantize(_CENTAVOS, rounding=ROUND_HALF_UP)
        modelo_valido = not erros and bool(grupos_disponiveis)
        avisos.extend(erros)
        if total_desativadas:
            avisos.append(
                f"{total_desativadas} conta(s) desativada(s) nao participarao da consulta."
            )

        return {
            "ok": True,
            "msg": (
                "Modelo confirmado e pronto para consultar."
                if modelo_valido else
                "O modelo foi lido, mas possui diferencas que precisam ser corrigidas."
            ),
            "caminho": str(arquivo),
            "arquivo": arquivo.name,
            "abas": abas,
            "abas_mensais": [nome for nome in abas if _ABAS_MENSAIS_RE.fullmatch(nome)],
            "aba_modelo": aba,
            "periodo": periodo,
            "modelo": configuracao,
            "modelo_id": configuracao.get("id") or "",
            "modelo_nome": configuracao["nome"],
            "modelo_valido": modelo_valido,
            "pronto_para_automacao": modelo_valido,
            "integracao_sap_disponivel": True,
            "resumo": {
                "contas_esperadas": len(contas_resumo),
                "contas_encontradas": len(contas_resumo),
                "contas_ativas": len(contas_resumo),
                "contas_desativadas": total_desativadas,
                "saldos_preenchidos": preenchidos,
                "saldos_vazios": len(saldos) - preenchidos,
                "saldos_zerados": zerados,
                "total_detalhes": _numero_front(total_atual),
                "total_verde": _numero_front(total_atual),
                "diferenca": None,
                "conferencia": "ok" if modelo_valido else "divergente",
            },
            "contas": contas_resumo,
            "grupos": grupos_disponiveis,
            "verificacoes": [],
            "avisos": list(dict.fromkeys(avisos)),
            "erros_modelo": list(dict.fromkeys(erros)),
        }
    finally:
        if workbook_formulas is not None:
            workbook_formulas.close()
        if workbook_valores is not None:
            workbook_valores.close()


def analisar_planilha_saldo_contas(
    caminho,
    aba_modelo: str = ABA_MODELO_SALDO_CONTAS,
    modelo=None,
) -> dict:
    if modelo is not None:
        return _analisar_planilha_com_modelo(caminho, modelo)

    arquivo = _validar_arquivo(caminho)
    workbook_formulas = None
    workbook_valores = None

    try:
        workbook_formulas = _carregar_planilha(arquivo, data_only=False)
        workbook_valores = _carregar_planilha(arquivo, data_only=True)

        abas = list(workbook_formulas.sheetnames)

        if aba_modelo not in workbook_formulas.sheetnames:
            return {
                "ok": False,
                "msg": f'A aba-modelo "{aba_modelo}" nao foi encontrada na planilha.',
                "caminho": str(arquivo),
                "arquivo": arquivo.name,
                "abas": abas,
            }

        sheet_formulas = workbook_formulas[aba_modelo]
        sheet_valores = workbook_valores[aba_modelo]
        cabecalho_contas = _normalizar_texto(sheet_formulas["C19"].value)
        cabecalho_valores = _normalizar_texto(sheet_formulas["D19"].value)
        periodo = _periodo_planilha(sheet_formulas["C1"].value)
        contas = []
        saldos_decimais = []
        contas_normalizadas = []

        for linha in range(LINHA_INICIAL_CONTAS, LINHA_FINAL_CONTAS + 1):
            conta = _normalizar_conta(sheet_formulas.cell(linha, 3).value)
            saldo = _decimal_centavos(sheet_valores.cell(linha, 4).value)
            contas_normalizadas.append(conta)
            saldos_decimais.append(saldo)
            contas.append(
                {
                    "linha": linha,
                    "conta": conta,
                    "saldo_atual": _numero_front(saldo),
                    "status": "vazio" if saldo is None else ("zero" if saldo == 0 else "preenchido"),
                }
            )

        grupos_disponiveis = ler_grupos_contas_modelo(sheet_formulas)
        contas_esperadas = list(CONTAS_APLICACAO_MODELO)
        contas_encontradas = [conta for conta in contas_normalizadas if conta]
        contas_ausentes = [conta for conta in contas_esperadas if conta not in contas_encontradas]
        contas_inesperadas = [conta for conta in contas_encontradas if conta not in contas_esperadas]
        contas_duplicadas = sorted(
            {conta for conta in contas_encontradas if contas_encontradas.count(conta) > 1}
        )
        lista_contas_ok = contas_normalizadas == contas_esperadas
        verificacoes = [
            _criar_verificacao(
                "aba_modelo",
                "Aba-modelo",
                True,
                f"{aba_modelo} localizada.",
            ),
            _criar_verificacao(
                "cabecalhos",
                "Cabecalhos do bloco",
                cabecalho_contas == "APLIC. MOVIMENTO" and cabecalho_valores == "INDIVIDUAIS",
                f"C19: {cabecalho_contas or '--'} | D19: {cabecalho_valores or '--'}.",
            ),
            _criar_verificacao(
                "lista_contas",
                "Lista de contas",
                lista_contas_ok and not contas_duplicadas,
                f"{len(contas_encontradas)} de {TOTAL_CONTAS_MODELO} contas do modelo.",
            ),
        ]

        formulas_ok = True

        for celula, formula_esperada in _FORMULAS_ESPERADAS.items():
            formula_atual = sheet_formulas[celula].value
            formula_ok = _normalizar_formula(formula_atual) == _normalizar_formula(formula_esperada)
            formulas_ok = formulas_ok and formula_ok
            verificacoes.append(
                _criar_verificacao(
                    f"formula_{celula.lower()}",
                    f"Formula {celula}",
                    formula_ok,
                    formula_atual or "Formula ausente.",
                )
            )

        total_detalhes = sum((saldo or Decimal("0")) for saldo in saldos_decimais).quantize(
            _CENTAVOS,
            rounding=ROUND_HALF_UP,
        )
        total_azul = _decimal_centavos(sheet_valores["C61"].value)

        if total_azul is None:
            total_azul = _decimal_centavos(sheet_formulas["C61"].value)

        if total_azul is None:
            total_azul = _valor_formula_constante(sheet_formulas["C61"].value)

        total_verde_cache = _decimal_centavos(sheet_valores["D61"].value)
        total_verde = total_verde_cache if total_verde_cache is not None else total_detalhes
        diferenca = None
        conferencia = "pendente"

        if total_azul is not None and total_verde is not None:
            diferenca = (total_azul - total_verde).quantize(_CENTAVOS, rounding=ROUND_HALF_UP)
            conferencia = "ok" if diferenca == 0 else "divergente"

        saldos_preenchidos = sum(1 for saldo in saldos_decimais if saldo is not None)
        saldos_zerados = sum(1 for saldo in saldos_decimais if saldo == 0)
        abas_mensais = [aba for aba in abas if _ABAS_MENSAIS_RE.fullmatch(aba)]
        avisos = []

        if contas_ausentes:
            avisos.append(f"Contas ausentes: {', '.join(contas_ausentes)}.")

        if contas_inesperadas:
            avisos.append(f"Contas diferentes do modelo: {', '.join(contas_inesperadas)}.")

        if contas_duplicadas:
            avisos.append(f"Contas duplicadas: {', '.join(contas_duplicadas)}.")

        if saldos_preenchidos < TOTAL_CONTAS_MODELO:
            avisos.append(
                f"A aba possui {TOTAL_CONTAS_MODELO - saldos_preenchidos} saldo(s) vazio(s). "
                "Vazio nao deve ser tratado como saldo zero."
            )

        if conferencia == "divergente":
            avisos.append("O total azul e a soma verde apresentam diferenca.")

        if not formulas_ok:
            avisos.append("Uma ou mais formulas de conferencia diferem do modelo SET.26.")

        modelo_valido = (
            cabecalho_contas == "APLIC. MOVIMENTO"
            and cabecalho_valores == "INDIVIDUAIS"
            and lista_contas_ok
            and not contas_duplicadas
            and formulas_ok
        )

        return {
            "ok": True,
            "msg": "Planilha analisada com sucesso.",
            "caminho": str(arquivo),
            "arquivo": arquivo.name,
            "abas": abas,
            "abas_mensais": abas_mensais,
            "aba_modelo": aba_modelo,
            "periodo": periodo,
            "modelo_valido": modelo_valido,
            "pronto_para_automacao": modelo_valido,
            "integracao_sap_disponivel": True,
            "bloco": {
                "cabecalho": "C19:D19",
                "contas": "C20:C60",
                "valores": "D20:D60",
                "total_independente": "C61",
                "total_formula": "D61",
            },
            "resumo": {
                "contas_esperadas": TOTAL_CONTAS_MODELO,
                "contas_encontradas": len(contas_encontradas),
                "saldos_preenchidos": saldos_preenchidos,
                "saldos_vazios": TOTAL_CONTAS_MODELO - saldos_preenchidos,
                "saldos_zerados": saldos_zerados,
                "total_detalhes": _numero_front(total_detalhes),
                "total_azul": _numero_front(total_azul),
                "total_verde": _numero_front(total_verde),
                "diferenca": _numero_front(diferenca),
                "conferencia": conferencia,
            },
            "contas": contas,
            "grupos": grupos_disponiveis,
            "verificacoes": verificacoes,
            "avisos": avisos,
        }
    finally:
        if workbook_formulas is not None:
            workbook_formulas.close()

        if workbook_valores is not None:
            workbook_valores.close()
